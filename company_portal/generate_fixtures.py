from pathlib import Path
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from docx import Document
from docx.shared import Pt, RGBColor

# fixture directory path
SAMPLE_DIR = Path(__file__).parent / "sample_files"


def create_apex_pdf_invoice(overwrite: bool = False):
    # generate apex pdf
    file_path = SAMPLE_DIR / "Apex_INV_2024_001.pdf"
    if file_path.exists() and not overwrite:
        return file_path
    SAMPLE_DIR.mkdir(parents=True, exist_ok=True)

    c = canvas.Canvas(str(file_path), pagesize=letter)
    width, height = letter

    # draw header bar
    c.setFillColor(colors.HexColor("#1e3a8a"))
    c.rect(0, height - 70, width, 70, fill=1, stroke=0)

    # draw company title
    c.setFillColor(colors.white)
    c.setFont("Helvetica-Bold", 20)
    c.drawString(40, height - 45, "APEX TECHNOLOGIES")

    # draw invoice details
    c.setFillColor(colors.HexColor("#0f172a"))
    c.setFont("Helvetica-Bold", 14)
    c.drawString(40, height - 110, "INVOICE")

    c.setFont("Helvetica-Bold", 10)
    c.drawString(40, height - 140, "INVOICE NUMBER:")
    c.setFont("Helvetica", 10)
    c.drawString(160, height - 140, "INV-2024-001")

    c.setFont("Helvetica-Bold", 10)
    c.drawString(40, height - 165, "PAYMENT DUE DATE:")
    c.setFont("Helvetica", 10)
    c.setFillColor(colors.HexColor("#b91c1c"))
    c.drawString(160, height - 165, "2026-11-15")

    # draw total amount
    c.setFillColor(colors.HexColor("#0f172a"))
    c.setFont("Helvetica-Bold", 12)
    c.drawString(40, height - 210, "TOTAL AMOUNT DUE:")
    c.setFont("Helvetica-Bold", 14)
    c.setFillColor(colors.HexColor("#1e3a8a"))
    c.drawString(180, height - 210, "$4,500.00")

    c.save()
    return file_path


def create_global_docx_invoice(overwrite: bool = False):
    # generate docx invoice
    file_path = SAMPLE_DIR / "GlobalLogistics_INV_8821.docx"
    if file_path.exists() and not overwrite:
        return file_path
    SAMPLE_DIR.mkdir(parents=True, exist_ok=True)

    doc = Document()

    # insert title paragraph
    p_title = doc.add_paragraph()
    run = p_title.add_run("GLOBAL LOGISTICS INC. - INVOICE")
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42)

    # insert details paragraphs
    doc.add_paragraph("Vendor: Global Logistics Inc.")
    doc.add_paragraph("Invoice Number: GL-INV-8821")
    doc.add_paragraph("Invoice Date: 2026-09-28")
    doc.add_paragraph("Payment Due Date: 2026-11-20")

    # insert total amount
    p_total = doc.add_paragraph()
    r = p_total.add_run("Total Amount: $12,850.75 USD")
    r.font.bold = True
    r.font.size = Pt(13)

    doc.save(str(file_path))
    return file_path


def generate_pdf_invoice(vendor_name: str, invoice_number: str, amount: float, due_date: str, filename: str = None) -> Path:
    # generate custom pdf
    SAMPLE_DIR.mkdir(parents=True, exist_ok=True)
    fname = filename or f"{vendor_name.replace(' ', '')}_{invoice_number.replace('-', '_')}.pdf"
    file_path = SAMPLE_DIR / fname

    c = canvas.Canvas(str(file_path), pagesize=letter)
    width, height = letter

    # draw header bar
    c.setFillColor(colors.HexColor("#0f172a"))
    c.rect(0, height - 70, width, 70, fill=1, stroke=0)

    # draw company title
    c.setFillColor(colors.white)
    c.setFont("Helvetica-Bold", 18)
    c.drawString(40, height - 45, vendor_name.upper())

    # draw invoice details
    c.setFillColor(colors.HexColor("#0f172a"))
    c.setFont("Helvetica-Bold", 14)
    c.drawString(40, height - 110, "INVOICE")

    c.setFont("Helvetica-Bold", 10)
    c.drawString(40, height - 140, "INVOICE NUMBER:")
    c.setFont("Helvetica", 10)
    c.drawString(160, height - 140, invoice_number)

    c.setFont("Helvetica-Bold", 10)
    c.drawString(40, height - 165, "PAYMENT DUE DATE:")
    c.setFont("Helvetica", 10)
    c.setFillColor(colors.HexColor("#b91c1c"))
    c.drawString(160, height - 165, due_date)

    # draw total amount
    c.setFillColor(colors.HexColor("#0f172a"))
    c.setFont("Helvetica-Bold", 12)
    c.drawString(40, height - 210, "TOTAL AMOUNT DUE:")
    c.setFont("Helvetica-Bold", 14)
    c.setFillColor(colors.HexColor("#1e3a8a"))
    c.drawString(180, height - 210, f"${amount:,.2f}")

    c.save()
    return file_path


def generate_docx_invoice(vendor_name: str, invoice_number: str, amount: float, due_date: str, filename: str = None) -> Path:
    # generate custom docx
    SAMPLE_DIR.mkdir(parents=True, exist_ok=True)
    fname = filename or f"{vendor_name.replace(' ', '')}_{invoice_number.replace('-', '_')}.docx"
    file_path = SAMPLE_DIR / fname

    doc = Document()
    p_title = doc.add_paragraph()
    run = p_title.add_run(f"{vendor_name.upper()} - INVOICE")
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph(f"Vendor: {vendor_name}")
    doc.add_paragraph(f"Invoice Number: {invoice_number}")
    doc.add_paragraph(f"Payment Due Date: {due_date}")

    p_total = doc.add_paragraph()
    r = p_total.add_run(f"Total Amount: ${amount:,.2f} USD")
    r.font.bold = True
    r.font.size = Pt(13)

    doc.save(str(file_path))
    return file_path


def create_all_fixtures(overwrite: bool = False):
    # generate all fixtures
    create_apex_pdf_invoice(overwrite=overwrite)
    create_global_docx_invoice(overwrite=overwrite)


if __name__ == "__main__":
    import argparse

    # parse cli arguments
    parser = argparse.ArgumentParser(description="CentrAlign Fixture Generator")
    parser.add_argument("--vendor", help="Vendor name")
    parser.add_argument("--invoice", default="INV-001", help="Invoice reference")
    parser.add_argument("--amount", type=float, default=1000.0, help="Amount in USD")
    parser.add_argument("--due", default="2026-12-31", help="Due date")
    parser.add_argument("--filename", default=None, help="Target filename")
    args = parser.parse_args()

    if args.vendor:
        # generate custom record
        p = generate_pdf_invoice(args.vendor, args.invoice, args.amount, args.due, args.filename)
        print(f"Created fixture: {p.name}")
    else:
        # run fixture generation
        create_all_fixtures()
        print("Generated all default enterprise fixtures.")
