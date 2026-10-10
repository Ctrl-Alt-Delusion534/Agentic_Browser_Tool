from pathlib import Path
import re
from typing import Dict, Any, Optional, List
from pypdf import PdfReader
from docx import Document
from tools.base import BaseTool, ToolResult


class DocumentParserTool(BaseTool):
    # tool metadata
    name = "parse_document"
    description = "Extracts text and financial fields from any PDF, DOCX, or text invoice."
    parameters_schema = {
        "type": "object",
        "properties": {
            "file_path": {
                "type": "string",
                "description": "Path to the invoice file to extract."
            }
        },
        "required": ["file_path"]
    }

    def _read_pdf(self, path: Path) -> str:
        # read pdf
        reader = PdfReader(str(path))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(pages)

    def _read_docx(self, path: Path) -> str:
        # read docx
        doc = Document(str(path))
        lines = [p.text for p in doc.paragraphs if p.text]
        for table in doc.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(lines)

    def _extract_vendor(self, lines: List[str]) -> Optional[str]:
        # labeled vendor
        for line in lines[:10]:
            match = re.search(r"^(?:vendor|from|biller|company\s*name)[:\s]+([A-Za-z0-9&.,\s'-]{3,50})", line, re.I)
            if match:
                return match.group(1).strip()

        # top header
        for line in lines[:6]:
            cleaned = line.strip()
            # skip boilerplate
            if not cleaned or re.search(r"^(invoice|bill\s+to|date|page|attn)", cleaned, re.I):
                continue
            if len(cleaned) >= 3 and not re.match(r"^[^a-zA-Z]+$", cleaned):
                return cleaned.split(" - ")[0].strip()

        return None

    def _extract_invoice_number(self, text: str) -> Optional[str]:
        # match patterns
        patterns = [
            r"(?i)(?:invoice|inv)\s*(?:number|no\.?|#|id|ref|code)\s*[:#]\s*([A-Za-z0-9-_/]+)",
            r"(?i)(?:invoice|inv)\s*(?:number|no\.?|#|id|ref|code)\s+([A-Za-z0-9-_/]*\d[A-Za-z0-9-_/]*)",
            r"(?i)\binvoice\s*:\s*([A-Za-z0-9-_/]+)",
            r"\b([A-Z]{2,5}-INV-\d+)\b",
            r"\b(INV-\d{4}-\d+)\b",
        ]
        for pat in patterns:
            match = re.search(pat, text)
            if match:
                val = match.group(1).strip()
                if not re.match(r"^(date|due|total|amount|vendor)$", val, re.I):
                    return val
        return None

    def _extract_amount(self, text: str) -> Optional[float]:
        # labeled total
        total_match = re.search(r"(?i)(?:total(?:\s+amount|\s+due|\s+payable)?|amount\s+due|grand\s+total)[:\s]*\$?\s*([0-9,]+\.[0-9]{2})", text)
        if total_match:
            return float(total_match.group(1).replace(",", ""))

        # currency amounts
        all_amounts = re.findall(r"\$\s*([0-9,]+\.[0-9]{2})", text)
        if all_amounts:
            # pick highest
            numbers = [float(a.replace(",", "")) for a in all_amounts]
            return max(numbers)

        return None

    def _extract_due_date(self, text: str) -> Optional[str]:
        # labeled due
        match = re.search(r"(?i)(?:payment\s+due|due\s+date|due\s+by)[:\s]*([0-9]{4}[-/][0-9]{2}[-/][0-9]{2})", text)
        if match:
            return match.group(1).replace("/", "-")

        # fallback date
        all_dates = re.findall(r"\b(\d{4}-\d{2}-\d{2})\b", text)
        if all_dates:
            return all_dates[-1]

        return None

    def execute(self, file_path: str) -> ToolResult:
        # validate path
        path = Path(file_path)
        if not path.exists():
            return ToolResult(success=False, error=f"File not found: {file_path}")

        # parse content
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            raw_text = self._read_pdf(path)
        elif suffix == ".docx":
            raw_text = self._read_docx(path)
        elif suffix in [".txt", ".csv"]:
            raw_text = path.read_text(encoding="utf-8", errors="ignore")
        else:
            return ToolResult(success=False, error=f"Unsupported format: {suffix}")

        # split lines
        lines = [line.strip() for line in raw_text.splitlines() if line.strip()]

        # extract fields
        data = {
            "file_name": path.name,
            "vendor_name": self._extract_vendor(lines),
            "invoice_number": self._extract_invoice_number(raw_text),
            "amount": self._extract_amount(raw_text),
            "due_date": self._extract_due_date(raw_text),
            "file_path": str(path.resolve()),
            "raw_text_preview": raw_text[:300].strip(),
        }

        # return output
        return ToolResult(success=True, data=data)
